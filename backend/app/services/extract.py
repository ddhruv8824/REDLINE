"""
app/services/extract.py
───────────────────────
Plain-text extraction from resume files (PDF via pdfplumber, DOCX via python-docx).
Used for providers that can't read PDFs natively.

Multi-column PDFs are the hard case: their text layer interleaves the columns line by
line ("Jenkins   Unison 2024-2025"), which scrambles jobs and skills together. So PDF
pages are extracted in layout mode and, when a column boundary is found, each band of
the page is read column by column (left column first, then right).
"""

from __future__ import annotations

from pathlib import Path

_MIN_TWO_COLUMN_LINES = 6   # need this many split lines before trusting a boundary
_GUTTER = 2                 # spaces that must separate the columns at the boundary


def _column_boundary(lines: list[str]) -> int | None:
    """Character position where a right-hand column starts, if the page has two columns."""
    width = max((len(l) for l in lines), default=0)
    if width < 40:
        return None
    best, best_score = None, 0
    for c in range(int(width * 0.15), int(width * 0.7)):
        split = crossing = 0
        for line in lines:
            if len(line) <= c:
                continue
            gutter_empty = line[c - _GUTTER:c].strip() == ""
            if gutter_empty and line[c] != " ":
                split += 1       # something starts exactly at the boundary
            elif line[c - 1] != " " and line[c] != " ":
                crossing += 1    # a word runs across the boundary
        score = split - 2 * crossing
        if split >= _MIN_TWO_COLUMN_LINES and score > best_score:
            best, best_score = c, score
    return best


def _reflow_columns(text: str) -> str:
    """Read a layout-mode page column by column instead of line by line."""
    lines = [l.rstrip() for l in text.splitlines()]
    c = _column_boundary([l for l in lines if l.strip()])
    if c is None:
        return "\n".join(l.strip() for l in lines if l.strip())

    out: list[str] = []
    left: list[tuple[int, str]] = []   # (line number, text)
    right: list[tuple[int, str]] = []

    def flush(carry_heading: bool = False) -> None:
        # Left-only lines after the right column's last line are headings for the
        # full-width text that follows (e.g. "PROJECTS" above a wide description),
        # so they go after the right column, not before it.
        last_right = right[-1][0] if right else -1
        carried = [t for n, t in left if n > last_right] if carry_heading else []
        out.extend(t for n, t in left if not (carry_heading and n > last_right))
        out.extend(t for _, t in right)
        out.extend(carried)
        left.clear()
        right.clear()

    for n, line in enumerate(lines):
        if not line.strip():
            continue
        crosses = len(line) > c and line[c - 1] != " " and line[c] != " "
        if crosses:          # full-width text (headline, summary, a wide section)
            flush(carry_heading=True)
            out.append(line.strip())
            continue
        if line[:c].strip():
            left.append((n, line[:c].strip()))
        if line[c:].strip():
            right.append((n, line[c:].strip()))
    flush()
    return "\n".join(out)


def extract_text_from_pdf(path: Path) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(layout=True, x_tolerance=3, y_tolerance=3)
            if text and text.strip():
                pages.append(_reflow_columns(text))
    return "\n\n".join(pages)


def extract_text_from_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in paragraphs:
                    paragraphs.append(text)
    return "\n".join(paragraphs)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext == ".docx":
        return extract_text_from_docx(path)
    raise ValueError(f"Unsupported file type: {ext}. Use PDF or DOCX.")
